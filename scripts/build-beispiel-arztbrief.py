#!/usr/bin/env python3
"""
Erzeugt die Downloads des Ratgebers „Arztbrief schreiben mit KI":
  public/ratgeber/beispiel-arztbrief-vorlage.docx / .pdf      (lesbare Vorlage)
  public/ratgeber/beispiel-arztbrief-platzhalter.docx / .pdf  (wie nach custix)

Quelle ist src/content/ratgeber/beispiel-arztbrief.json (auch die Website
liest daraus). Aufruf: python3 scripts/build-beispiel-arztbrief.py
Braucht python-docx und reportlab.

Die Platzhalter-Fassung ersetzt {{TYP:Wert}} mechanisch nach dem Schema von
custix ([PERSON_1], [DATUM_1], …). Vor dem Go-live soll sie durch einen echten
Durchlauf der Vorlage durch custix ersetzt werden (Spec mrgoofman/custix-ai#6).
"""
import json, re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_RIGHT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src/content/ratgeber/beispiel-arztbrief.json"
OUT = ROOT / "public/ratgeber"
MARKUP = re.compile(r"\{\{([A-Z]+):([^}]+)\}\}")

letter = json.loads(SRC.read_text(encoding="utf-8"))


class Numbering:
    """Gleicher Wert → gleicher Platzhalter, nummeriert in Lesereihenfolge."""

    def __init__(self):
        self.ids, self.counters = {}, {}

    def placeholder(self, typ, value):
        key = (typ, value)
        if key not in self.ids:
            self.counters[typ] = self.counters.get(typ, 0) + 1
            self.ids[key] = f"[{typ}_{self.counters[typ]}]"
        return self.ids[key]


def segments(text, variant, numbering):
    """Liste aus (text, is_placeholder)."""
    out, last = [], 0
    for m in MARKUP.finditer(text):
        if m.start() > last:
            out.append((text[last:m.start()], False))
        typ, value = m.group(1), m.group(2)
        out.append((numbering.placeholder(typ, value), True) if variant == "placeholders" else (value, False))
        last = m.end()
    if last < len(text):
        out.append((text[last:], False))
    return out


def plain(segs):
    return "".join(t for t, _ in segs)


# ---------------------------------------------------------------- Word
def build_docx(variant, path):
    numbering = Numbering()
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin = s.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    def para(segs, bold=False, size=None, align=None, space_after=6, italic=False, color=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if align == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for text, is_ph in segs:
            r = p.add_run(text)
            r.bold = bold or is_ph
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor(*color)
            if is_ph:
                r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
                r.font.highlight_color = 7  # WD_COLOR_INDEX.YELLOW
        return p

    para([(letter["notice"], False)], size=8, italic=True, color=(0x94, 0xA3, 0xB8), space_after=12)
    for block in letter["blocks"]:
        kind = block["kind"]
        if kind == "header":
            for i, line in enumerate(block["lines"]):
                para(segments(line, variant, numbering), bold=(i == 0), size=12 if i == 0 else 10, space_after=0)
            doc.add_paragraph()
        elif kind == "recipient":
            for line in block["lines"]:
                para(segments(line, variant, numbering), space_after=0)
            doc.add_paragraph()
        elif kind == "date":
            para(segments(block["text"], variant, numbering), align="right", space_after=12)
        elif kind == "subject":
            para(segments(block["text"], variant, numbering), bold=True, space_after=12)
        elif kind in ("salutation", "paragraph"):
            para(segments(block["text"], variant, numbering), space_after=10)
        elif kind == "section":
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            r = h.add_run(block["heading"])
            r.bold = True
            r.font.size = Pt(11.5)
            r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
            for line in block["lines"]:
                para(segments(line, variant, numbering), space_after=4)
        elif kind == "closing":
            doc.add_paragraph()
            for i, line in enumerate(block["lines"]):
                para(segments(line, variant, numbering), space_after=0 if i else 14)
    doc.save(path)


# ---------------------------------------------------------------- PDF
def esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def rich(segs):
    return "".join(
        f'<b><font backColor="#FEF3C7" color="#1E3A5F">{esc(t)}</font></b>' if ph else esc(t)
        for t, ph in segs
    )


def build_pdf(variant, path):
    numbering = Numbering()
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=10.5, leading=14.5, spaceAfter=5)
    st = {
        "notice": ParagraphStyle("notice", parent=base, fontName="Helvetica-Oblique", fontSize=8, leading=10, textColor="#94A3B8", spaceAfter=12),
        "org": ParagraphStyle("org", parent=base, fontName="Helvetica-Bold", fontSize=12, leading=15, spaceAfter=0),
        "tight": ParagraphStyle("tight", parent=base, spaceAfter=0),
        "date": ParagraphStyle("date", parent=base, alignment=TA_RIGHT, spaceAfter=12),
        "subject": ParagraphStyle("subject", parent=base, fontName="Helvetica-Bold", spaceAfter=12),
        "heading": ParagraphStyle("heading", parent=base, fontName="Helvetica-Bold", fontSize=11, textColor="#1E3A5F", spaceBefore=8, spaceAfter=3),
    }
    story = [Paragraph(esc(letter["notice"]), st["notice"])]
    for block in letter["blocks"]:
        kind = block["kind"]
        if kind == "header":
            for i, line in enumerate(block["lines"]):
                story.append(Paragraph(rich(segments(line, variant, numbering)), st["org"] if i == 0 else st["tight"]))
            story.append(Spacer(1, 6 * mm))
        elif kind == "recipient":
            for line in block["lines"]:
                story.append(Paragraph(rich(segments(line, variant, numbering)), st["tight"]))
            story.append(Spacer(1, 6 * mm))
        elif kind == "date":
            story.append(Paragraph(rich(segments(block["text"], variant, numbering)), st["date"]))
        elif kind == "subject":
            story.append(Paragraph(rich(segments(block["text"], variant, numbering)), st["subject"]))
        elif kind in ("salutation", "paragraph"):
            story.append(Paragraph(rich(segments(block["text"], variant, numbering)), base))
        elif kind == "section":
            story.append(Paragraph(esc(block["heading"]), st["heading"]))
            for line in block["lines"]:
                story.append(Paragraph(rich(segments(line, variant, numbering)), base))
        elif kind == "closing":
            story.append(Spacer(1, 6 * mm))
            for i, line in enumerate(block["lines"]):
                story.append(Paragraph(rich(segments(line, variant, numbering)), base if i == 0 else st["tight"]))
                if i == 0:
                    story.append(Spacer(1, 8 * mm))
    SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=20 * mm, bottomMargin=20 * mm,
        title=letter["title"], author="custix.ai", subject="Beispieldaten – erfunden",
    ).build(story)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for variant, name in (("template", "vorlage"), ("placeholders", "platzhalter")):
        build_docx(variant, OUT / f"beispiel-arztbrief-{name}.docx")
        build_pdf(variant, OUT / f"beispiel-arztbrief-{name}.pdf")
    for f in sorted(OUT.iterdir()):
        print(f"{f.relative_to(ROOT)}  {f.stat().st_size // 1024} KB")
